import docx

def extract_tables(cell, prefix=""):
    for inner_table in cell.tables:
        for i, r in enumerate(inner_table.rows):
            cells_text = []
            for c in r.cells:
                text = c.text.replace('\n', ' ').replace('\r', ' ').strip()
                cells_text.append(text)
            print(f"{prefix}Row {i}: | " + " | ".join(cells_text) + " |")

def inspect_table():
    doc = docx.Document("Basic_Kanji_320_Stroke_Order_Meanings_Examples.docx")
    
    with open("table_nested_inspect.txt", "w", encoding="utf-8") as f:
        import sys
        sys.stdout = f
        
        for t_idx, table in enumerate(doc.tables):
            print(f"--- Main Table {t_idx} ---")
            for i, row in enumerate(table.rows):
                print(f"Row {i}:")
                for j, cell in enumerate(row.cells):
                    text = cell.text.replace('\n', ' ').replace('\r', ' ').strip()
                    print(f"  Col {j} text: {text}")
                    extract_tables(cell, prefix=f"  Col {j} Inner Table -> ")

if __name__ == '__main__':
    inspect_table()
