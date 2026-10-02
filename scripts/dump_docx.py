import docx

def dump_docx(filename, output_name):
    try:
        doc = docx.Document(filename)
        with open(output_name, 'w', encoding='utf-8') as f:
            for para in doc.paragraphs:
                if para.text.strip():
                    f.write(para.text + '\n')
            
            f.write('\n--- TABLES ---\n')
            for i, table in enumerate(doc.tables):
                f.write(f'\nTable {i+1}\n')
                for row in table.rows:
                    row_text = [cell.text.replace('\n', ' ').strip() for cell in row.cells]
                    f.write(' | '.join(row_text) + '\n')
                    
        print(f"Dumped {filename} to {output_name}")
    except Exception as e:
        print(f"Error reading {filename}: {e}")

if __name__ == '__main__':
    dump_docx("Minna_no_Nihongo_I_Grammar_by_Lesson.docx", "grammar_dump.txt")
    dump_docx("Basic_Kanji_320_Stroke_Order_Meanings_Examples.docx", "kanji_dump.txt")
